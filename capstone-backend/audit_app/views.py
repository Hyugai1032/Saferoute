import datetime
import ipaddress
import json
from io import BytesIO

from django.apps import apps
from django.conf import settings
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework.exceptions import PermissionDenied
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from . import conf
from .models import AuditLog
from .serializers import AuditLogSerializer


def _municipality_names(ids):
    ids = {i for i in ids if i is not None}
    if not ids:
        return {}
    try:
        Municipality = apps.get_model("auth_app", "Municipality")
    except LookupError:
        return {}
    return dict(Municipality.objects.filter(pk__in=ids).values_list("pk", "name"))


def _day_start(day):
    value = datetime.datetime.combine(day, datetime.time.min)
    return timezone.make_aware(value) if settings.USE_TZ else value


class AuditLogListView(APIView):
    """
    GET /audit-logs/
    Filters: action, resource_type, resource_id, actor (user id), request_id,
             q (email, name, record label, path or IP), date_from, date_to (YYYY-MM-DD)
    Paging:  page_size (max 100) and cursor (the next_cursor from the previous page).

    Uses keyset ("load more") paging on the id instead of page numbers, so page 500
    is as fast as page 1 however large the table grows. The total count is only
    computed for the first page.

    Provincial admins see everything; municipal admins see their own municipality.
    """
    permission_classes = [IsAuthenticated]

    def get_filtered_queryset(self, request):
        user = request.user
        role = getattr(user, "role", None)
        if not (user.is_superuser or role in conf.viewer_roles()):
            raise PermissionDenied("You are not allowed to view the system logs.")

        qs = AuditLog.objects.all()
        if not (user.is_superuser or role in conf.all_scope_roles()):
            mun_id = getattr(user, "municipality_id", None)
            qs = qs.filter(municipality_id=mun_id) if mun_id else qs.none()

        p = request.query_params

        action = (p.get("action") or "").upper()
        if action and action != "ALL":
            qs = qs.filter(action=action)

        # success = 2xx/3xx, failed = 4xx/5xx
        outcome = (p.get("outcome") or "").lower()
        if outcome == "success":
            qs = qs.filter(status_code__lt=400)
        elif outcome == "failed":
            qs = qs.filter(status_code__gte=400)

        resource_type = p.get("resource_type") or ""
        if resource_type and resource_type.lower() != "all":
            qs = qs.filter(resource_type=resource_type)
            if p.get("resource_id"):
                qs = qs.filter(resource_id=p["resource_id"])

        if p.get("actor"):
            if not p["actor"].isdigit():
                return None, Response({"actor": "Must be a user id."}, status=400)
            qs = qs.filter(actor_id=int(p["actor"]))

        if p.get("request_id"):
            qs = qs.filter(request_id=p["request_id"][:32])

        # Plain ranges on the indexed column (no DATE() call, so MySQL can use the index).
        for key in ("date_from", "date_to"):
            raw = p.get(key)
            if not raw:
                continue
            day = parse_date(raw)
            if not day:
                return None, Response({key: "Use the format YYYY-MM-DD."}, status=400)
            if key == "date_from":
                qs = qs.filter(timestamp__gte=_day_start(day))
            else:
                qs = qs.filter(timestamp__lt=_day_start(day + datetime.timedelta(days=1)))

        q = (p.get("q") or "").strip()[:100]
        if q:
            cond = (
                Q(actor_email__icontains=q) | Q(actor_name__icontains=q)
                | Q(resource_label__icontains=q) | Q(path__icontains=q)
            )
            if _looks_like_ip(q):
                cond |= Q(ip_address=q)
            if q.lstrip("#").isdigit():
                cond |= Q(resource_id=q.lstrip("#"))
            qs = qs.filter(cond)

        return qs, None

    def get(self, request):
        qs, error = self.get_filtered_queryset(request)
        if error:
            return error

        p = request.query_params
        try:
            page_size = min(max(int(p.get("page_size", 25)), 1), 100)
            cursor = int(p["cursor"]) if p.get("cursor") else None
        except ValueError:
            return Response({"detail": "page_size and cursor must be integers."}, status=400)

        count = qs.count() if cursor is None else None
        page_qs = qs.filter(id__lt=cursor) if cursor else qs
        rows = list(page_qs.order_by("-id")[:page_size + 1])
        has_more = len(rows) > page_size
        rows = rows[:page_size]

        names = _municipality_names(r.municipality_id for r in rows)
        return Response({
            "count": count,
            "next_cursor": rows[-1].id if has_more and rows else None,
            "results": AuditLogSerializer(rows, many=True, context={"municipality_names": names}).data,
        })


def _looks_like_ip(text):
    try:
        ipaddress.ip_address(text)
        return True
    except ValueError:
        return False


def format_details(action, details):
    """Readable one-cell version of `details` for the Excel export."""
    if not details:
        return ""
    if not isinstance(details, dict):
        return json.dumps(details, default=str)
    lines = []
    for key, value in details.items():
        if action == AuditLog.Action.UPDATE and isinstance(value, list) and len(value) == 2:
            lines.append(f"{key}: {_cell(value[0])} → {_cell(value[1])}")
        else:
            lines.append(f"{key}: {_cell(value)}")
    return "\n".join(lines)


def _cell(value):
    if value is None or value == "":
        return "(empty)"
    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str)
    return str(value)


class AuditLogExportView(AuditLogListView):
    """GET /audit-logs/export/ - every entry matching the filters, as .xlsx (max 10,000)."""
    MAX_ROWS = 10000

    COLUMNS = [
        ("Date & Time", 20), ("Action", 18), ("User", 26), ("User Email", 28), ("Role", 22),
        ("Municipality", 18), ("Record Type", 18), ("Record ID", 10), ("Record", 32),
        ("Details", 60), ("Method", 9), ("Path", 36), ("Status", 8), ("IP Address", 16),
        ("Request ID", 34),
    ]

    def get(self, request):
        qs, error = self.get_filtered_queryset(request)
        if error:
            return error

        count = qs.count()
        if count > self.MAX_ROWS:
            return Response(
                {"detail": f"Too many entries to export ({count:,}). "
                           f"Narrow the date range or filters to {self.MAX_ROWS:,} or fewer."},
                status=400,
            )

        try:
            from openpyxl import Workbook
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response({"detail": "Excel export needs openpyxl on the server."}, status=501)

        def clean(value):
            return ILLEGAL_CHARACTERS_RE.sub("", "" if value is None else str(value))

        def naive(dt):
            if dt is None:
                return None
            return timezone.localtime(dt).replace(tzinfo=None) if timezone.is_aware(dt) else dt

        rows = list(qs.order_by("-id"))
        names = _municipality_names(r.municipality_id for r in rows)
        role_labels = {}
        try:
            role_labels = dict(apps.get_model("auth_app", "CustomUser").ROLE_CHOICES)
        except (LookupError, AttributeError):
            pass

        wb = Workbook()
        ws = wb.active
        ws.title = "System Logs"
        for col, (title, width) in enumerate(self.COLUMNS, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[1].height = 22

        for r, log in enumerate(rows, start=2):
            ts = naive(log.timestamp)
            values = [
                ts, log.get_action_display(), log.actor_name or ("System" if not log.method else "Guest"),
                log.actor_email, role_labels.get(log.actor_role, log.actor_role),
                names.get(log.municipality_id, ""), log.resource_type, log.resource_id,
                log.resource_label, format_details(log.action, log.details), log.method,
                log.path, log.status_code, log.ip_address, log.request_id,
            ]
            for c, value in enumerate(values, start=1):
                if isinstance(value, str) or value is None:
                    cell = ws.cell(row=r, column=c)
                    cell.value = clean(value)
                    cell.data_type = "s"  # never evaluate "=..." as a formula
                else:
                    cell = ws.cell(row=r, column=c, value=value)
                if c == 1 and ts:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss AM/PM"
                wrap = c in (9, 10, 12)
                cell.alignment = Alignment(wrap_text=wrap, vertical="top")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(self.COLUMNS))}{max(ws.max_row, 1)}"

        buffer = BytesIO()
        wb.save(buffer)
        stamp = naive(timezone.now()).date().isoformat()
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="System_Logs_{stamp}.xlsx"'
        return response


class AuditLogMetaView(APIView):
    """GET /audit-logs/meta/ - choices for the page's filter dropdowns (no database query)."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        if not (user.is_superuser or getattr(user, "role", None) in conf.viewer_roles()):
            raise PermissionDenied("You are not allowed to view the system logs.")

        types = []
        for label in conf.tracked_models():
            try:
                types.append(apps.get_model(label).__name__)
            except (LookupError, ValueError):
                continue
        return Response({
            "actions": [{"value": v, "label": l} for v, l in AuditLog.Action.choices],
            "resource_types": sorted(set(types)),
        })
